import json
from io import BytesIO
from urllib.parse import quote

from collections import defaultdict

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from .. import models, row_match_cache, schemas
from ..connectors import (
    ConnectionError_,
    build_extract_alternatives,
    create_connector,
    fuzzy_index_keys,
    values_equal_with_extracts,
    values_fuzzy_equal,
    values_sig_hash,
    values_sig_hashes_with_extracts,
    _stringify_row,
    _stringify_value,
)
from ..database import get_db

router = APIRouter(prefix="/api/compare", tags=["compare"])


def _get_task(db: Session, task_id: int) -> models.Task:
    task = db.query(models.Task).filter(models.Task.id == task_id).first()
    if not task:
        raise HTTPException(status_code=404, detail="任务不存在")
    return task


def _connector(task: models.Task, side: str):
    if side not in ("source", "target"):
        raise HTTPException(status_code=400, detail="side 只能是 source 或 target")
    raw = task.source_config if side == "source" else task.target_config
    try:
        return create_connector(json.loads(raw))
    except ConnectionError_ as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


def _run(fn):
    try:
        return fn()
    except ConnectionError_ as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"查询失败: {exc}") from exc


def _column_source_fields(item: schemas.ColumnPair) -> list[str]:
    """Physical source columns for a match pair; supports N→1 aggregation via '+'."""
    fields = [str(f).strip() for f in (item.source_fields or []) if str(f).strip()]
    if fields:
        return fields
    name = str(item.source or "").strip()
    if "+" in name:
        parts = [p.strip() for p in name.split("+") if p.strip()]
        if parts:
            return parts
    return [name] if name else []


def _join_composite_cell_values(values: list) -> str:
    parts: list[str] = []
    for value in values:
        if value is None:
            parts.append("NULL")
            continue
        text = str(value)
        parts.append("NULL" if text == "" else text)
    return "+".join(parts)


def _row_source_match_value(row: dict, item: schemas.ColumnPair):
    fields = _column_source_fields(item)
    if not fields:
        return None
    if len(fields) == 1:
        return row.get(fields[0])
    return _join_composite_cell_values([row.get(field) for field in fields])


def _joined_source_match_value(
    resolved_multi: dict[str, list[dict]],
    table: str,
    item: schemas.ColumnPair,
):
    fields = _column_source_fields(item)
    if not fields:
        return None
    rows = resolved_multi.get(table) or []
    if len(fields) == 1:
        return _aggregate_join_column(rows, fields[0])
    return _join_composite_cell_values([_aggregate_join_column(rows, field) for field in fields])


def _physical_source_columns(
    columns: list[schemas.ColumnPair],
    primary: str | None = None,
    only_table: str | None = None,
) -> list[str]:
    """Collect physical source columns; optionally restrict to one source table."""
    out: list[str] = []
    seen: set[str] = set()
    for item in columns:
        table = item.source_table or primary or ""
        if only_table is not None and table != only_table:
            continue
        for field in _column_source_fields(item):
            if field in seen:
                continue
            seen.add(field)
            out.append(field)
    return out


def _has_composite_source(columns: list[schemas.ColumnPair]) -> bool:
    return any(len(_column_source_fields(item)) > 1 for item in columns)


@router.get("/{task_id}/databases")
def list_databases(task_id: int, side: str = Query(...), db: Session = Depends(get_db)):
    task = _get_task(db, task_id)
    connector = _connector(task, side)
    names = _run(connector.list_databases)
    return {"side": side, "databases": names}


@router.get("/{task_id}/schemas")
def list_schemas(
    task_id: int,
    side: str = Query(...),
    database: str = Query(...),
    db: Session = Depends(get_db),
):
    task = _get_task(db, task_id)
    connector = _connector(task, side)
    names = _run(lambda: connector.list_schemas(database))
    return {"side": side, "database": database, "schemas": names}


@router.get("/{task_id}/tables")
def list_tables(
    task_id: int,
    side: str = Query(...),
    database: str = Query(...),
    schema_name: str | None = Query(default=None),
    db: Session = Depends(get_db),
):
    task = _get_task(db, task_id)
    connector = _connector(task, side)
    tables = _run(lambda: connector.list_tables(database, schema_name))
    return {"side": side, "database": database, "schema_name": schema_name, "tables": tables}


@router.get("/{task_id}/columns")
def list_columns(
    task_id: int,
    side: str = Query(...),
    database: str = Query(...),
    table: str = Query(...),
    schema_name: str | None = Query(default=None),
    db: Session = Depends(get_db),
):
    task = _get_task(db, task_id)
    connector = _connector(task, side)
    columns = _run(lambda: connector.list_columns(database, table, schema_name))
    return {
        "side": side,
        "database": database,
        "table": table,
        "schema_name": schema_name,
        "columns": columns,
        "count": len(columns),
    }


@router.get("/{task_id}/rows")
def list_rows(
    task_id: int,
    side: str = Query(...),
    database: str = Query(...),
    table: str = Query(...),
    schema_name: str | None = Query(default=None),
    limit: int = Query(default=1000, ge=1, le=10000),
    offset: int = Query(default=0, ge=0),
    db: Session = Depends(get_db),
):
    task = _get_task(db, task_id)
    connector = _connector(task, side)
    data = _run(lambda: connector.fetch_rows(database, table, schema_name, limit, offset))
    return {
        "side": side,
        "database": database,
        "table": table,
        "schema_name": schema_name,
        **data,
    }


@router.get("/{task_id}/row-count")
def row_count(
    task_id: int,
    side: str = Query(...),
    database: str = Query(...),
    table: str = Query(...),
    schema_name: str | None = Query(default=None),
    db: Session = Depends(get_db),
):
    task = _get_task(db, task_id)
    connector = _connector(task, side)
    total = _run(lambda: connector.count_rows(database, table, schema_name))
    return {"side": side, "database": database, "table": table, "total": total}


@router.get("/{task_id}/values")
def distinct_values(
    task_id: int,
    side: str = Query(...),
    database: str = Query(...),
    table: str = Query(...),
    column: str = Query(...),
    schema_name: str | None = Query(default=None),
    limit: int = Query(default=200, ge=1, le=1000),
    db: Session = Depends(get_db),
):
    task = _get_task(db, task_id)
    connector = _connector(task, side)
    values = _run(lambda: connector.distinct_values(database, table, column, schema_name, limit))
    return {
        "side": side,
        "database": database,
        "table": table,
        "column": column,
        "values": values,
        "count": len(values),
    }


def _pair_key(payload: schemas.MatchPairRef) -> str:
    return row_match_cache.pair_key(
        payload.source_database,
        payload.source_table,
        payload.source_schema,
        payload.target_database,
        payload.target_table,
        payload.target_schema,
    )


def _get_cache(task_id: int, payload: schemas.MatchPairRef) -> row_match_cache.RowMatchCache | None:
    return row_match_cache.get(task_id, _pair_key(payload))


def _norm_alias_value(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _build_value_maps(aliases: list[schemas.ValueAlias]) -> tuple[dict, dict]:
    source_map: dict[tuple[str, str], str] = {}
    target_map: dict[tuple[str, str], str] = {}
    for alias in aliases:
        token = f"{alias.target_field}={_norm_alias_value(alias.target_value)}"
        source_map[(alias.source_field, _norm_alias_value(alias.source_value))] = token
        target_map[(alias.target_field, _norm_alias_value(alias.target_value))] = token
    return source_map, target_map


def _join_partner_specs(
    primary: str, joins: list[schemas.SourceJoinSpec]
) -> list[tuple[str, str, str]]:
    """Return [(partner_table, local_field_on_primary, partner_field), ...] for direct hops."""
    specs = []
    for join in joins:
        if join.left_table == primary:
            specs.append((join.right_table, join.left_field, join.right_field))
        elif join.right_table == primary:
            specs.append((join.left_table, join.right_field, join.left_field))
    return specs


def _join_edge_from(table: str, join: schemas.SourceJoinSpec) -> tuple[str, str, str] | None:
    """Return (local_field, remote_table, remote_field) when table is on either side."""
    if join.left_table == table:
        return join.left_field, join.right_table, join.right_field
    if join.right_table == table:
        return join.right_field, join.left_table, join.left_field
    return None


def _fetch_partner_rows(
    connector, database: str, schema_name: str | None, table: str, field: str, keys: list
) -> list[dict]:
    """Load all partner rows for keys, retrying when LIMIT truncates a batch."""
    if not keys:
        return []
    cap = max(len(keys) * 50, 5000)
    rows = connector.fetch_rows_by_column_values(
        database, table, field, keys, schema_name, limit=cap
    )
    if len(rows) < cap or len(keys) <= 1:
        return rows
    found = {str(row.get(field)) for row in rows if row.get(field) is not None}
    missing = [key for key in keys if str(key) not in found]
    if not missing:
        return rows
    mid = max(1, len(keys) // 2)
    return _fetch_partner_rows(
        connector, database, schema_name, table, field, keys[:mid]
    ) + _fetch_partner_rows(
        connector, database, schema_name, table, field, keys[mid:]
    )


def _index_partner_rows(connector, database: str, schema_name: str | None, table: str, field: str, keys: list):
    index: dict[str, dict] = {}
    uniq = []
    seen = set()
    for key in keys:
        if key is None:
            continue
        token = str(key)
        if token in seen:
            continue
        seen.add(token)
        uniq.append(key)
    chunk = 2000
    for i in range(0, len(uniq), chunk):
        part = uniq[i : i + chunk]
        for row in _fetch_partner_rows(connector, database, schema_name, table, field, part):
            token = str(row.get(field)) if row.get(field) is not None else None
            if token is not None and token not in index:
                index[token] = row
    return index


def _load_join_partner_indexes(
    connector,
    database: str,
    schema_name: str | None,
    primary: str,
    primary_rows: list[dict],
    joins: list[schemas.SourceJoinSpec],
) -> dict[str, dict[str, dict]]:
    """BFS-load partner indexes so bridge tables can reach multi-hop sources."""
    partner_indexes: dict[str, dict[str, dict]] = {}
    frontier: list[tuple[str, list[dict]]] = [(primary, primary_rows)]
    loaded = {primary}

    while frontier:
        next_frontier: list[tuple[str, list[dict]]] = []
        for table, rows in frontier:
            for join in joins:
                edge = _join_edge_from(table, join)
                if not edge:
                    continue
                local_field, remote_table, remote_field = edge
                if remote_table in loaded:
                    continue
                keys = [row.get(local_field) for row in rows]
                index = _index_partner_rows(
                    connector, database, schema_name, remote_table, remote_field, keys
                )
                partner_indexes[remote_table] = index
                loaded.add(remote_table)
                next_frontier.append((remote_table, list(index.values())))
        frontier = next_frontier
    return partner_indexes


def _join_edges_between(
    table: str, remote_table: str, joins: list[schemas.SourceJoinSpec]
) -> list[tuple[str, str]]:
    """Return [(local_field, remote_field), ...] from table toward remote_table."""
    edges: list[tuple[str, str]] = []
    for join in joins or []:
        edge = _join_edge_from(table, join)
        if not edge:
            continue
        local_field, rem_table, remote_field = edge
        if rem_table == remote_table:
            edges.append((local_field, remote_field))
    return edges


def _values_equal_for_join(left, right) -> bool:
    if left is None or right is None:
        return left is None and right is None
    return str(left) == str(right)


def _partner_matches_all_edges(local_row: dict, partner: dict, edges: list[tuple[str, str]]) -> bool:
    """Multi-field joins between the same tables are AND (intersection)."""
    for local_field, remote_field in edges or []:
        if not _values_equal_for_join(local_row.get(local_field), partner.get(remote_field)):
            return False
    return True


def _resolve_joined_row(
    primary: str,
    primary_row: dict,
    joins: list[schemas.SourceJoinSpec],
    partner_indexes: dict[str, dict[str, dict]],
) -> dict[str, dict | None]:
    """Map every reachable table (incl. primary) to its row for this primary row."""
    resolved: dict[str, dict | None] = {primary: primary_row}
    queue = [primary]
    while queue:
        table = queue.pop(0)
        row = resolved.get(table)
        if not row:
            continue
        for join in joins:
            edge = _join_edge_from(table, join)
            if not edge:
                continue
            local_field, remote_table, remote_field = edge
            if remote_table in resolved:
                continue
            all_edges = _join_edges_between(table, remote_table, joins) or [
                (local_field, remote_field)
            ]
            key = row.get(local_field)
            token = str(key) if key is not None else None
            partner = None
            if token is not None:
                candidate = partner_indexes.get(remote_table, {}).get(token)
                if candidate is not None and _partner_matches_all_edges(row, candidate, all_edges):
                    partner = candidate
            resolved[remote_table] = partner
            if partner is not None:
                queue.append(remote_table)
    return resolved


def _normalize_match_conditions(raw_conditions) -> list[dict]:
    normalized = []
    for item in raw_conditions or []:
        data = item.model_dump() if hasattr(item, "model_dump") else dict(item)
        terms = [str(t).strip() for t in (data.get("keywords") or []) if str(t).strip()]
        if not terms:
            continue
        mode = str(data.get("mode") or "").strip().lower()
        if mode not in {"include", "exclude", "equals", "not_equals"}:
            mode = "exclude" if data.get("negate") else "include"
        normalized.append(
            {
                "search_columns": [c for c in (data.get("search_columns") or []) if c],
                "keywords": terms,
                "mode": mode,
                "negate": mode in {"exclude", "not_equals"},
                "logic": _normalize_condition_logic(data.get("logic")),
            }
        )
    return normalized


def _condition_columns(conditions: list[dict]) -> list[str]:
    cols = []
    for cond in conditions:
        cols.extend(cond.get("search_columns") or [])
    return list(dict.fromkeys(cols))


def _parse_column_ref(col: str, default_table: str) -> tuple[str, str]:
    text = str(col or "").strip()
    if "::" in text:
        table, field = text.split("::", 1)
        return (table or default_table), field
    return default_table, text


def _primary_condition_fields(conditions: list[dict], primary: str) -> list[str]:
    fields = []
    for col in _condition_columns(conditions):
        table, field = _parse_column_ref(col, primary)
        if table == primary and field:
            fields.append(field)
    return list(dict.fromkeys(fields))


def _conditions_need_join(conditions: list[dict], primary: str) -> bool:
    for col in _condition_columns(conditions):
        table, _ = _parse_column_ref(col, primary)
        if table != primary:
            return True
    return False


def _stringify_search_value(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, bytes):
        return value.hex()
    text = str(value)
    return None if text.upper() == "NULL" else text


def _is_null_search_keyword(term) -> bool:
    return str(term or "").strip().upper() == "NULL"


def _is_empty_search_text(text: str | None) -> bool:
    return text is None or not str(text).strip()


def _row_search_text(row: dict, columns: list[str]) -> str:
    parts = []
    for column in columns:
        value = _stringify_search_value(row.get(column))
        if _is_empty_search_text(value):
            # 空值在界面显示为 NULL；写入 NULL 便于关键词 "NULL" 模糊命中
            parts.append(f"{column}=NULL")
            parts.append("NULL")
            continue
        parts.append(f"{column}={value}")
        parts.append(str(value))
    return " ".join(parts)


def _fuzzy_match_text(text: str, keyword: str) -> bool:
    import re

    kw = (keyword or "").strip().lower()
    if not kw:
        return True
    haystack = (text or "").lower()
    compact = re.sub(r"[\s_\-./]", "", haystack)
    for part in kw.split():
        if not part:
            continue
        if part in haystack:
            continue
        compact_part = re.sub(r"[\s_\-./]", "", part)
        if compact_part and compact_part in compact:
            continue
        return False
    return True


def _exact_match_values(row: dict, columns: list[str], keywords: list[str]) -> bool:
    targets = {str(term).strip() for term in keywords if str(term).strip()}
    if not targets:
        return True
    null_wanted = any(_is_null_search_keyword(term) for term in targets)
    plain_targets = {term for term in targets if not _is_null_search_keyword(term)}
    for column in columns:
        value = _stringify_search_value(row.get(column))
        text = "" if value is None else str(value).strip()
        if _is_empty_search_text(text):
            if null_wanted:
                return True
            continue
        if text in plain_targets:
            return True
    return False


def _normalize_condition_logic(logic) -> str:
    value = str(logic or "and").strip().lower()
    return "or" if value == "or" else "and"


def _combine_condition_hits(
    hits: list[bool],
    logics: list[str] | None = None,
    default_logic: str = "and",
) -> bool:
    if not hits:
        return True
    fallback = _normalize_condition_logic(default_logic)
    if not logics or len(logics) != len(hits):
        return any(hits) if fallback == "or" else all(hits)
    result = bool(hits[0])
    for index in range(1, len(hits)):
        op = _normalize_condition_logic(logics[index] if index < len(logics) else fallback)
        if op == "or":
            result = result or bool(hits[index])
        else:
            result = result and bool(hits[index])
    return result


def _condition_hit(
    *,
    mode: str,
    keywords: list[str],
    exact_hit: bool,
    fuzzy_hit: bool,
) -> bool:
    if mode in {"equals", "not_equals"}:
        hit = exact_hit
        if mode == "not_equals":
            hit = not hit
        return bool(hit)
    hit = fuzzy_hit
    if mode == "exclude":
        hit = not hit
    return bool(hit)


def _row_matches_conditions(
    row: dict, conditions: list[dict], default_cols: list[str], logic: str = "and"
) -> bool:
    if not conditions:
        return True
    hits = []
    logics = []
    for cond in conditions:
        match_cols = cond["search_columns"] or default_cols
        mode = cond.get("mode") or ("exclude" if cond.get("negate") else "include")
        if mode in {"equals", "not_equals"}:
            exact_hit = _exact_match_values(row, match_cols, cond["keywords"])
            hit = _condition_hit(mode=mode, keywords=cond["keywords"], exact_hit=exact_hit, fuzzy_hit=False)
        else:
            haystack = _row_search_text(row, match_cols)
            fuzzy_hit = any(_fuzzy_match_text(haystack, term) for term in cond["keywords"])
            hit = _condition_hit(mode=mode, keywords=cond["keywords"], exact_hit=False, fuzzy_hit=fuzzy_hit)
        hits.append(hit)
        logics.append(cond.get("logic") or logic)
    return _combine_condition_hits(hits, logics, logic)


def _index_partner_rows_multi(connector, database: str, schema_name: str | None, table: str, field: str, keys: list):
    index: dict[str, list[dict]] = {}
    uniq = []
    seen = set()
    for key in keys:
        if key is None:
            continue
        token = str(key)
        if token in seen:
            continue
        seen.add(token)
        uniq.append(key)
    chunk = 2000
    for i in range(0, len(uniq), chunk):
        part = uniq[i : i + chunk]
        for row in _fetch_partner_rows(connector, database, schema_name, table, field, part):
            token = str(row.get(field)) if row.get(field) is not None else None
            if token is None:
                continue
            index.setdefault(token, []).append(row)
    return index


def _load_join_partner_indexes_multi(
    connector,
    database: str,
    schema_name: str | None,
    primary: str,
    primary_rows: list[dict],
    joins: list[schemas.SourceJoinSpec],
) -> dict[str, dict[str, list[dict]]]:
    partner_indexes: dict[str, dict[str, list[dict]]] = {}
    frontier: list[tuple[str, list[dict]]] = [(primary, primary_rows)]
    loaded = {primary}

    while frontier:
        next_frontier: list[tuple[str, list[dict]]] = []
        for table, rows in frontier:
            for join in joins:
                edge = _join_edge_from(table, join)
                if not edge:
                    continue
                local_field, remote_table, remote_field = edge
                if remote_table in loaded:
                    continue
                keys = [row.get(local_field) for row in rows]
                index = _index_partner_rows_multi(
                    connector, database, schema_name, remote_table, remote_field, keys
                )
                partner_indexes[remote_table] = index
                loaded.add(remote_table)
                next_frontier.append((remote_table, [row for group in index.values() for row in group]))
        frontier = next_frontier
    return partner_indexes


def _resolve_joined_rows_multi(
    primary: str,
    primary_row: dict,
    joins: list[schemas.SourceJoinSpec],
    partner_indexes: dict[str, dict[str, list[dict]]],
) -> dict[str, list[dict]]:
    resolved: dict[str, list[dict]] = {primary: [primary_row]}
    queue = [primary]
    while queue:
        table = queue.pop(0)
        rows = resolved.get(table) or []
        if not rows:
            continue
        for join in joins:
            edge = _join_edge_from(table, join)
            if not edge:
                continue
            local_field, remote_table, remote_field = edge
            if remote_table in resolved:
                continue
            all_edges = _join_edges_between(table, remote_table, joins) or [
                (local_field, remote_field)
            ]
            partners: list[dict] = []
            seen = set()
            for row in rows:
                key = row.get(local_field)
                token = str(key) if key is not None else None
                if token is None:
                    continue
                for partner in partner_indexes.get(remote_table, {}).get(token, []):
                    if not _partner_matches_all_edges(row, partner, all_edges):
                        continue
                    identity = id(partner)
                    if identity in seen:
                        continue
                    seen.add(identity)
                    partners.append(partner)
            resolved[remote_table] = partners
            if partners:
                queue.append(remote_table)
    return resolved


def _aggregate_join_column(rows: list[dict], col: str):
    """Join-field compare value: one row keeps native value; many rows join like the UI."""
    if not rows:
        return None
    values = []
    seen = set()
    for row in rows:
        raw = row.get(col) if row else None
        token = "" if raw is None else str(raw).strip()
        if not token or token.upper() == "NULL":
            continue
        if token in seen:
            continue
        seen.add(token)
        values.append(raw)
    if not values:
        return None
    if len(values) == 1:
        return values[0]
    return "，".join(str(v).strip() for v in values)


def _resolved_search_text(resolved: dict[str, list[dict]], columns: list[str], default_table: str) -> str:
    parts = []
    for column in columns:
        table, field = _parse_column_ref(column, default_table)
        rows = resolved.get(table) or []
        label = f"{table}.{field}" if table != default_table else field
        if not rows:
            parts.append(f"{label}=NULL")
            parts.append("NULL")
            continue
        for row in rows:
            value = _stringify_search_value(row.get(field))
            if _is_empty_search_text(value):
                parts.append(f"{label}=NULL")
                parts.append("NULL")
                continue
            parts.append(f"{label}={value}")
            parts.append(str(value))
    return " ".join(parts)


def _resolved_exact_match(resolved: dict[str, list[dict]], columns: list[str], keywords: list[str], default_table: str) -> bool:
    targets = {str(term).strip() for term in keywords if str(term).strip()}
    if not targets:
        return True
    null_wanted = any(_is_null_search_keyword(term) for term in targets)
    plain_targets = {term for term in targets if not _is_null_search_keyword(term)}
    for column in columns:
        table, field = _parse_column_ref(column, default_table)
        rows = resolved.get(table) or [{}]
        for row in rows:
            value = _stringify_search_value(row.get(field))
            text = "" if value is None else str(value).strip()
            if _is_empty_search_text(text):
                if null_wanted:
                    return True
                continue
            if text in plain_targets:
                return True
    return False


def _resolved_matches_conditions(
    resolved: dict[str, list[dict]],
    conditions: list[dict],
    default_table: str,
    default_cols: list[str],
    logic: str = "and",
) -> bool:
    if not conditions:
        return True
    hits = []
    logics = []
    for cond in conditions:
        match_cols = cond["search_columns"] or default_cols
        mode = cond.get("mode") or ("exclude" if cond.get("negate") else "include")
        if mode in {"equals", "not_equals"}:
            exact_hit = _resolved_exact_match(resolved, match_cols, cond["keywords"], default_table)
            hit = _condition_hit(mode=mode, keywords=cond["keywords"], exact_hit=exact_hit, fuzzy_hit=False)
        else:
            haystack = _resolved_search_text(resolved, match_cols, default_table)
            fuzzy_hit = any(_fuzzy_match_text(haystack, term) for term in cond["keywords"])
            hit = _condition_hit(mode=mode, keywords=cond["keywords"], exact_hit=False, fuzzy_hit=fuzzy_hit)
        hits.append(hit)
        logics.append(cond.get("logic") or logic)
    return _combine_condition_hits(hits, logics, logic)


def _search_rows_with_joins(
    connector,
    database: str,
    primary: str,
    schema_name: str | None,
    columns: list[str],
    conditions: list[dict],
    joins: list[schemas.SourceJoinSpec],
    offset: int,
    limit: int,
    allowed_indices: set[int] | None = None,
    condition_logic: str = "and",
) -> dict:
    partner_specs = _join_partner_specs(primary, joins)
    # columns 可能混入关联表字段名；主表 SELECT 只保留关联键与主表条件列
    primary_fields = set(_primary_condition_fields(conditions, primary))
    for _, local_field, _ in partner_specs:
        if local_field:
            primary_fields.add(local_field)
    for join in joins:
        if join.left_table == primary and join.left_field:
            primary_fields.add(join.left_field)
        if join.right_table == primary and join.right_field:
            primary_fields.add(join.right_field)
    # 仅当列名不含 "." / "::" 时才视为主表列（前端跨表字段用 table::field）
    for col in columns or []:
        name = str(col or "").strip()
        if not name or "::" in name or "." in name:
            continue
        primary_fields.add(name)
    primary_cols = [c for c in primary_fields if c]
    if not primary_cols:
        return {"rows": [], "indices": [], "total": 0}

    primary_rows = list(connector.iter_column_rows(database, primary, primary_cols, schema_name))
    partner_indexes = _load_join_partner_indexes_multi(
        connector, database, schema_name, primary, primary_rows, joins
    )

    matched = 0
    page_rows: list[dict] = []
    page_indices: list[int] = []
    end = offset + limit
    default_cols = columns or primary_cols
    for idx, primary_row in enumerate(primary_rows):
        if allowed_indices is not None and idx not in allowed_indices:
            continue
        resolved = _resolve_joined_rows_multi(primary, primary_row, joins, partner_indexes)
        if not _resolved_matches_conditions(
            resolved, conditions, primary, default_cols, condition_logic
        ):
            continue
        if offset <= matched < end:
            page_rows.append(_stringify_row(primary_row))
            page_indices.append(idx)
        matched += 1
    return {"rows": page_rows, "indices": page_indices, "total": matched}


@router.post("/{task_id}/row-match")
def match_rows(task_id: int, payload: schemas.RowMatchRequest, db: Session = Depends(get_db)):
    task = _get_task(db, task_id)
    source = _connector(task, "source")
    target = _connector(task, "target")
    primary = payload.source_table
    match_columns = list(payload.columns or [])
    source_tables = [item.source_table or primary for item in match_columns]
    has_foreign_sources = any(t != primary for t in source_tables)
    needs_join = has_foreign_sources and bool(payload.source_joins)
    if has_foreign_sources and not needs_join:
        # 未配置源表关联时，忽略其它源表字段，避免把 duty_catalog.code 查到 resource
        match_columns = [
            item for item in match_columns if (item.source_table or primary) == primary
        ]
        source_tables = [item.source_table or primary for item in match_columns]
    source_cols = [item.source for item in match_columns]
    target_cols = [item.target for item in match_columns]
    if not source_cols or not target_cols:
        raise HTTPException(status_code=400, detail="请先连线要对比的字段")
    source_map, target_map = _build_value_maps(payload.value_aliases or [])
    extract_alternatives = build_extract_alternatives(
        source_cols,
        target_cols,
        payload.extract_match_rules or [],
    )
    has_extracts = bool(extract_alternatives)
    first_pair_extracts = extract_alternatives.get((source_cols[0], target_cols[0]))
    has_composites = _has_composite_source(match_columns)
    # 非 join 路径只能查主表物理列，避免把其它源表字段误查到主表
    physical_source_cols = _physical_source_columns(
        match_columns, primary=primary, only_table=primary
    )
    fuzzy = bool(payload.fuzzy_content_match)
    source_conditions = _normalize_match_conditions(payload.source_conditions)
    target_conditions = _normalize_match_conditions(payload.target_conditions)
    source_logic = _normalize_condition_logic(payload.source_condition_logic)
    target_logic = _normalize_condition_logic(payload.target_condition_logic)
    has_source_filter = bool(source_conditions)
    has_target_filter = bool(target_conditions)
    filtered = has_source_filter or has_target_filter

    def _build_exact():
        buckets: dict[bytes, list[int]] = defaultdict(list)
        target_universe: list[int] | None = [] if has_target_filter else None
        target_total = 0
        target_value_rows: dict[int, list] | None = {} if has_extracts else None

        if has_extracts or has_target_filter:
            fetch_cols = list(dict.fromkeys([*target_cols, *_condition_columns(target_conditions)]))
            for index, row in enumerate(
                target.iter_column_rows(
                    payload.target_database,
                    payload.target_table,
                    fetch_cols,
                    payload.target_schema,
                )
            ):
                if has_target_filter and not _row_matches_conditions(
                    row, target_conditions, fetch_cols, target_logic
                ):
                    continue
                if target_universe is not None:
                    target_universe.append(index)
                values = [row.get(col) for col in target_cols]
                if target_value_rows is not None:
                    target_value_rows[index] = values
                if has_extracts:
                    for signature in values_sig_hashes_with_extracts(
                        target_cols,
                        values,
                        source_cols,
                        target_map,
                        extract_alternatives,
                        "target",
                    ):
                        buckets[signature].append(index)
                else:
                    buckets[values_sig_hash(target_cols, values, target_map)].append(index)
                target_total += 1
        else:
            for index, signature in enumerate(
                target.iter_signature_hashes(
                    payload.target_database,
                    payload.target_table,
                    target_cols,
                    payload.target_schema,
                    target_map,
                )
            ):
                buckets[signature].append(index)
                target_total = index + 1

        source_to_target: dict[int, int] = {}
        target_to_source: dict[int, int] = {}
        used: set[int] = set()
        source_universe: list[int] | None = [] if has_source_filter else None
        source_total = 0

        def _accept_source(source_index: int, values: list | None, signature: bytes | None = None) -> None:
            nonlocal source_total
            if source_universe is not None:
                source_universe.append(source_index)
            source_total += 1
            if has_extracts:
                assert values is not None and target_value_rows is not None
                seen_candidates: set[int] = set()
                for sig in values_sig_hashes_with_extracts(
                    source_cols,
                    values,
                    target_cols,
                    source_map,
                    extract_alternatives,
                    "source",
                ):
                    for target_index in buckets.get(sig, ()):
                        if target_index in used or target_index in seen_candidates:
                            continue
                        seen_candidates.add(target_index)
                        if not values_equal_with_extracts(
                            source_cols,
                            values,
                            source_map,
                            target_cols,
                            target_value_rows[target_index],
                            target_map,
                            extract_alternatives,
                        ):
                            continue
                        used.add(target_index)
                        source_to_target[source_index] = target_index
                        target_to_source[target_index] = source_index
                        return
                return
            assert signature is not None
            for target_index in buckets.get(signature, ()):
                if target_index in used:
                    continue
                used.add(target_index)
                source_to_target[source_index] = target_index
                target_to_source[target_index] = source_index
                break

        if not needs_join:
            if has_extracts or has_source_filter or has_composites:
                fetch_cols = list(
                    dict.fromkeys([*physical_source_cols, *_condition_columns(source_conditions)])
                )
                for source_index, row in enumerate(
                    source.iter_column_rows(
                        payload.source_database,
                        payload.source_table,
                        fetch_cols,
                        payload.source_schema,
                    )
                ):
                    if has_source_filter and not _row_matches_conditions(
                        row, source_conditions, fetch_cols, source_logic
                    ):
                        continue
                    values = [_row_source_match_value(row, item) for item in match_columns]
                    if has_extracts:
                        _accept_source(source_index, values)
                    else:
                        _accept_source(
                            source_index,
                            None,
                            values_sig_hash(source_cols, values, source_map),
                        )
            else:
                for source_index, signature in enumerate(
                    source.iter_signature_hashes(
                        payload.source_database,
                        payload.source_table,
                        source_cols,
                        payload.source_schema,
                        source_map,
                    )
                ):
                    _accept_source(source_index, None, signature)
        else:
            joins = payload.source_joins or []
            partner_specs = _join_partner_specs(primary, joins)
            primary_fields = set()
            for item, table in zip(match_columns, source_tables):
                if table == primary:
                    primary_fields.update(_column_source_fields(item))
            for _, local_field, _ in partner_specs:
                primary_fields.add(local_field)
            if has_source_filter:
                primary_fields.update(_primary_condition_fields(source_conditions, primary))
            primary_cols = list(primary_fields)

            primary_rows = list(
                source.iter_column_rows(
                    payload.source_database,
                    payload.source_table,
                    primary_cols,
                    payload.source_schema,
                )
            )
            partner_indexes_multi = _load_join_partner_indexes_multi(
                source,
                payload.source_database,
                payload.source_schema,
                primary,
                primary_rows,
                joins,
            )
            for source_index, primary_row in enumerate(primary_rows):
                resolved_multi = _resolve_joined_rows_multi(
                    primary, primary_row, joins, partner_indexes_multi
                )
                if has_source_filter:
                    if _conditions_need_join(source_conditions, primary):
                        if not _resolved_matches_conditions(
                            resolved_multi,
                            source_conditions,
                            primary,
                            primary_cols,
                            source_logic,
                        ):
                            continue
                    elif not _row_matches_conditions(
                        primary_row, source_conditions, primary_cols, source_logic
                    ):
                        continue
                values = [
                    _joined_source_match_value(resolved_multi, table, item)
                    for item, table in zip(match_columns, source_tables)
                ]
                if has_extracts:
                    _accept_source(source_index, values)
                else:
                    _accept_source(
                        source_index,
                        None,
                        values_sig_hash(source_cols, values, source_map),
                    )

        return row_match_cache.RowMatchCache(
            pair_key=_pair_key(payload),
            source_to_target=source_to_target,
            target_to_source=target_to_source,
            source_total=source_total,
            target_total=target_total,
            source_universe=source_universe,
            target_universe=target_universe,
        )

    def _build_fuzzy():
        target_fetch_cols = list(dict.fromkeys([*target_cols, *_condition_columns(target_conditions)]))
        target_rows = list(
            target.iter_column_rows(
                payload.target_database,
                payload.target_table,
                target_fetch_cols,
                payload.target_schema,
            )
        )
        target_universe: list[int] | None = [] if has_target_filter else None
        buckets: dict[str, list[int]] = defaultdict(list)
        target_total = 0
        for index, row in enumerate(target_rows):
            if has_target_filter and not _row_matches_conditions(
                row, target_conditions, target_fetch_cols, target_logic
            ):
                continue
            if target_universe is not None:
                target_universe.append(index)
            target_total += 1
            for key in fuzzy_index_keys(
                target_cols[0],
                row.get(target_cols[0]),
                target_map,
                first_pair_extracts,
                "target",
            ):
                buckets[key].append(index)

        source_to_target: dict[int, int] = {}
        target_to_source: dict[int, int] = {}
        used: set[int] = set()
        source_universe: list[int] | None = [] if has_source_filter else None
        source_total = 0

        def _try_match(source_index: int, values: list) -> None:
            nonlocal source_total
            if source_universe is not None:
                source_universe.append(source_index)
            source_total += 1
            seen_candidates: set[int] = set()
            for key in fuzzy_index_keys(
                source_cols[0],
                values[0],
                source_map,
                first_pair_extracts,
                "source",
            ):
                for target_index in buckets.get(key, ()):
                    if target_index in used or target_index in seen_candidates:
                        continue
                    seen_candidates.add(target_index)
                    if not values_fuzzy_equal(
                        source_cols,
                        values,
                        source_map,
                        target_cols,
                        target_rows[target_index],
                        target_map,
                        extract_alternatives,
                    ):
                        continue
                    used.add(target_index)
                    source_to_target[source_index] = target_index
                    target_to_source[target_index] = source_index
                    return

        if not needs_join:
            source_fetch_cols = list(
                dict.fromkeys([*physical_source_cols, *_condition_columns(source_conditions)])
            )
            for source_index, row in enumerate(
                source.iter_column_rows(
                    payload.source_database,
                    payload.source_table,
                    source_fetch_cols,
                    payload.source_schema,
                )
            ):
                if has_source_filter and not _row_matches_conditions(
                    row, source_conditions, source_fetch_cols, source_logic
                ):
                    continue
                values = [_row_source_match_value(row, item) for item in match_columns]
                _try_match(source_index, values)
        else:
            joins = payload.source_joins or []
            partner_specs = _join_partner_specs(primary, joins)
            primary_fields = set()
            for item, table in zip(match_columns, source_tables):
                if table == primary:
                    primary_fields.update(_column_source_fields(item))
            for _, local_field, _ in partner_specs:
                primary_fields.add(local_field)
            if has_source_filter:
                primary_fields.update(_primary_condition_fields(source_conditions, primary))
            primary_cols = list(primary_fields)
            primary_rows = list(
                source.iter_column_rows(
                    payload.source_database,
                    payload.source_table,
                    primary_cols,
                    payload.source_schema,
                )
            )
            partner_indexes_multi = _load_join_partner_indexes_multi(
                source,
                payload.source_database,
                payload.source_schema,
                primary,
                primary_rows,
                joins,
            )
            for source_index, primary_row in enumerate(primary_rows):
                resolved_multi = _resolve_joined_rows_multi(
                    primary, primary_row, joins, partner_indexes_multi
                )
                if has_source_filter:
                    if _conditions_need_join(source_conditions, primary):
                        if not _resolved_matches_conditions(
                            resolved_multi,
                            source_conditions,
                            primary,
                            primary_cols,
                            source_logic,
                        ):
                            continue
                    elif not _row_matches_conditions(
                        primary_row, source_conditions, primary_cols, source_logic
                    ):
                        continue
                values = [
                    _joined_source_match_value(resolved_multi, table, item)
                    for item, table in zip(match_columns, source_tables)
                ]
                _try_match(source_index, values)

        return row_match_cache.RowMatchCache(
            pair_key=_pair_key(payload),
            source_to_target=source_to_target,
            target_to_source=target_to_source,
            source_total=source_total,
            target_total=target_total,
            source_universe=source_universe,
            target_universe=target_universe,
        )

    cache = _run(_build_fuzzy if fuzzy else _build_exact)
    row_match_cache.put(task_id, cache)
    return {
        "matched": cache.matched,
        "left_only": cache.source_total - cache.matched,
        "right_only": cache.target_total - cache.matched,
        "source_total": cache.source_total,
        "target_total": cache.target_total,
        "fuzzy_content_match": fuzzy,
        "filtered": filtered,
    }


@router.post("/{task_id}/row-match-overlay")
def row_match_overlay(task_id: int, payload: schemas.RowMatchOverlayRequest, db: Session = Depends(get_db)):
    _get_task(db, task_id)
    if len(payload.source_indices) > 2000 or len(payload.target_indices) > 2000:
        raise HTTPException(status_code=400, detail="单次最多查询 2000 个行号")
    cache = _get_cache(task_id, payload)
    if not cache:
        return {
            "has_match": False,
            "mappings": [],
            "mapped_source": [],
            "mapped_target": [],
        }
    return cache.overlay(payload.source_indices, payload.target_indices)


@router.post("/{task_id}/unmatched-rows")
def unmatched_rows(task_id: int, payload: schemas.UnmatchedRowsRequest, db: Session = Depends(get_db)):
    if payload.side not in ("source", "target"):
        raise HTTPException(status_code=400, detail="side 只能是 source 或 target")
    if payload.limit < 1 or payload.limit > 1000:
        raise HTTPException(status_code=400, detail="每页最多 1000 行")
    if payload.offset < 0:
        raise HTTPException(status_code=400, detail="offset 不能为负数")
    task = _get_task(db, task_id)
    cache = _get_cache(task_id, payload)
    if not cache:
        raise HTTPException(status_code=400, detail="请先进行内容对比")
    indices = cache.unmatched(payload.side)
    page = indices[payload.offset : payload.offset + payload.limit]
    connector = _connector(task, payload.side)
    rows = _run(
        lambda: connector.fetch_rows_by_indices(
            payload.database,
            payload.table,
            page,
            payload.schema_name,
        )
    )
    return {
        "side": payload.side,
        "rows": rows,
        "indices": page,
        "total": len(indices),
        "count": len(rows),
    }


@router.post("/{task_id}/search-rows")
def search_rows(task_id: int, payload: schemas.SearchRowsRequest, db: Session = Depends(get_db)):
    if payload.side not in ("source", "target"):
        raise HTTPException(status_code=400, detail="side 只能是 source 或 target")
    if payload.limit < 1 or payload.limit > 1000:
        raise HTTPException(status_code=400, detail="每页最多 1000 行")
    if payload.offset < 0:
        raise HTTPException(status_code=400, detail="offset 不能为负数")
    keyword = (payload.keyword or "").strip()
    keywords = [str(item).strip() for item in (payload.keywords or []) if str(item).strip()]
    if keyword and keyword not in keywords:
        keywords.append(keyword)
    conditions = []
    for item in payload.conditions or []:
        data = item.model_dump() if hasattr(item, "model_dump") else dict(item)
        terms = [str(t).strip() for t in (data.get("keywords") or []) if str(t).strip()]
        if not terms:
            continue
        mode = str(data.get("mode") or "").strip().lower()
        if mode not in {"include", "exclude", "equals", "not_equals"}:
            mode = "exclude" if data.get("negate") else "include"
        conditions.append(
            {
                "search_columns": [c for c in (data.get("search_columns") or []) if c],
                "keywords": terms,
                "mode": mode,
                "negate": mode in {"exclude", "not_equals"},
                "logic": _normalize_condition_logic(data.get("logic")),
            }
        )
    if not conditions and not keywords:
        raise HTTPException(status_code=400, detail="请输入搜索内容")
    if not payload.columns and not payload.search_columns and not any(c["search_columns"] for c in conditions):
        raise HTTPException(status_code=400, detail="请先连线要对比的字段")
    task = _get_task(db, task_id)
    allowed = None
    if payload.unmatched_only:
        cache = _get_cache(task_id, payload)
        if not cache:
            raise HTTPException(status_code=400, detail="请先进行内容对比")
        allowed = set(cache.unmatched(payload.side))
    connector = _connector(task, payload.side)
    joins = payload.source_joins or []
    use_join_search = (
        payload.side == "source"
        and bool(joins)
        and (
            _conditions_need_join(conditions, payload.table)
            or any("::" in str(c) for c in (payload.search_columns or []))
        )
    )
    condition_logic = _normalize_condition_logic(payload.condition_logic)
    if use_join_search:
        if not conditions and keywords:
            conditions = [
                {
                    "search_columns": [c for c in (payload.search_columns or payload.columns or []) if c],
                    "keywords": keywords,
                    "mode": "exclude" if payload.negate else "include",
                    "negate": bool(payload.negate),
                }
            ]
        result = _run(
            lambda: _search_rows_with_joins(
                connector,
                payload.database,
                payload.table,
                payload.schema_name,
                payload.columns,
                conditions,
                joins,
                payload.offset,
                payload.limit,
                allowed,
                condition_logic,
            )
        )
    else:
        result = _run(
            lambda: connector.search_rows(
                payload.database,
                payload.table,
                payload.columns,
                keyword,
                payload.schema_name,
                payload.offset,
                payload.limit,
                allowed,
                payload.search_columns,
                keywords,
                bool(payload.negate),
                conditions,
                condition_logic,
            )
        )
    return {
        "side": payload.side,
        **result,
    }


@router.post("/{task_id}/row-match-edit")
def row_match_edit(task_id: int, payload: schemas.RowMatchEditRequest, db: Session = Depends(get_db)):
    _get_task(db, task_id)
    action = (payload.action or "").lower()
    if action == "clear":
        cache = _get_cache(task_id, payload)
        if cache:
            row_match_cache.clear(task_id)
        return {"ok": True, "has_match": False}

    cache = _get_cache(task_id, payload)
    if not cache:
        raise HTTPException(status_code=400, detail="请先进行内容对比")
    if action == "link":
        if payload.source is None or payload.target is None:
            raise HTTPException(status_code=400, detail="连线需要源行号和目标行号")
        cache.link(payload.source, payload.target)
    elif action == "unlink":
        cache.unlink(payload.source, payload.target)
    else:
        raise HTTPException(status_code=400, detail="action 只能是 link / unlink / clear")
    return {
        "ok": True,
        "has_match": True,
        "matched": cache.matched,
        "left_only": cache.source_total - cache.matched,
        "right_only": cache.target_total - cache.matched,
        "source_total": cache.source_total,
        "target_total": cache.target_total,
    }


@router.post("/{task_id}/rows-by-indices")
def rows_by_indices(task_id: int, payload: schemas.RowsByIndicesRequest, db: Session = Depends(get_db)):
    if payload.side not in ("source", "target"):
        raise HTTPException(status_code=400, detail="side 只能是 source 或 target")
    if len(payload.indices) > 10000:
        raise HTTPException(status_code=400, detail="单次最多按 10000 个行号查询")
    task = _get_task(db, task_id)
    connector = _connector(task, payload.side)
    rows = _run(
        lambda: connector.fetch_rows_by_indices(
            payload.database,
            payload.table,
            payload.indices,
            payload.schema_name,
        )
    )
    return {
        "side": payload.side,
        "database": payload.database,
        "table": payload.table,
        "rows": rows,
        "count": len(rows),
    }


@router.post("/{task_id}/rows-by-values")
def rows_by_values(task_id: int, payload: schemas.RowsByColumnValuesRequest, db: Session = Depends(get_db)):
    if payload.side not in ("source", "target"):
        raise HTTPException(status_code=400, detail="side 只能是 source 或 target")
    if len(payload.values) > 5000:
        raise HTTPException(status_code=400, detail="单次最多按 5000 个值查询")
    limit = min(max(int(payload.limit or 2000), 1), 20000)
    task = _get_task(db, task_id)
    connector = _connector(task, payload.side)
    rows = _run(
        lambda: connector.fetch_rows_by_column_values(
            payload.database,
            payload.table,
            payload.column,
            payload.values,
            payload.schema_name,
            limit,
        )
    )
    return {
        "side": payload.side,
        "database": payload.database,
        "table": payload.table,
        "column": payload.column,
        "rows": rows,
        "count": len(rows),
    }


def _export_cell_text(value) -> str:
    if value is None:
        return ""
    text = _stringify_value(value)
    if text is None:
        return ""
    return str(text)


def _export_column_header(col: str, default_table: str) -> str:
    table, field = _parse_column_ref(col, default_table)
    if table and table != default_table:
        return f"{table}.{field}"
    return field or str(col)


def _resolved_export_value(resolved: dict, column: str, default_table: str):
    table, field = _parse_column_ref(column, default_table)
    rows = resolved.get(table) or []
    if not rows:
        return None
    if len(rows) == 1:
        return rows[0].get(field) if rows[0] else None
    return _aggregate_join_column(rows, field)


def _collect_export_side_rows(
    *,
    connector,
    database: str,
    table: str,
    schema_name: str | None,
    columns: list[str],
    conditions: list[dict],
    condition_logic: str,
    joins: list | None,
    allowed_indices: set[int] | None,
    limit: int,
) -> list[dict]:
    """Return list of {header: value} dicts for selected columns under filters."""
    if not columns:
        return []
    headers = [_export_column_header(col, table) for col in columns]
    use_join = bool(joins) and (
        _conditions_need_join(conditions, table)
        or any(("::" in str(c) or "." in str(c)) for c in columns)
    )
    out: list[dict] = []
    if use_join:
        partner_specs = _join_partner_specs(table, joins or [])
        primary_fields = set(_primary_condition_fields(conditions, table))
        for _, local_field, _ in partner_specs:
            if local_field:
                primary_fields.add(local_field)
        for join in joins or []:
            if join.left_table == table and join.left_field:
                primary_fields.add(join.left_field)
            if join.right_table == table and join.right_field:
                primary_fields.add(join.right_field)
        for col in columns:
            tname, field = _parse_column_ref(col, table)
            if tname == table and field:
                primary_fields.add(field)
        primary_cols = [c for c in primary_fields if c]
        if not primary_cols:
            return []
        primary_rows = list(connector.iter_column_rows(database, table, primary_cols, schema_name))
        partner_indexes = _load_join_partner_indexes_multi(
            connector, database, schema_name, table, primary_rows, joins or []
        )
        for idx, primary_row in enumerate(primary_rows):
            if allowed_indices is not None and idx not in allowed_indices:
                continue
            resolved = _resolve_joined_rows_multi(table, primary_row, joins or [], partner_indexes)
            if conditions and not _resolved_matches_conditions(
                resolved, conditions, table, columns, condition_logic
            ):
                continue
            row_out = {}
            for col, header in zip(columns, headers):
                row_out[header] = _export_cell_text(_resolved_export_value(resolved, col, table))
            out.append(row_out)
            if len(out) >= limit:
                break
        return out

    # no join: stream primary table with optional conditions
    fetch_cols = list(dict.fromkeys([*_primary_condition_fields(conditions, table), *columns]))
    fetch_cols = [c for c in fetch_cols if c and "::" not in c and "." not in c]
    if not fetch_cols:
        fetch_cols = [c for c in columns if c and "::" not in c and "." not in c]
    if not fetch_cols:
        return []
    for idx, row in enumerate(connector.iter_column_rows(database, table, fetch_cols, schema_name)):
        if allowed_indices is not None and idx not in allowed_indices:
            continue
        if conditions and not _row_matches_conditions(row, conditions, fetch_cols, condition_logic):
            continue
        row_out = {}
        for col, header in zip(columns, headers):
            _table, field = _parse_column_ref(col, table)
            row_out[header] = _export_cell_text(row.get(field))
        out.append(row_out)
        if len(out) >= limit:
            break
    return out


def _build_export_workbook(source_rows: list[dict], target_rows: list[dict], source_headers: list[str], target_headers: list[str]):
    from openpyxl import Workbook

    wb = Workbook(write_only=True)
    if source_headers:
        ws = wb.create_sheet("源数据表字段内容")
        ws.append(source_headers)
        for row in source_rows:
            ws.append([row.get(h, "") for h in source_headers])
    if target_headers:
        ws = wb.create_sheet("目标数据表字段内容")
        ws.append(target_headers)
        for row in target_rows:
            ws.append([row.get(h, "") for h in target_headers])
    if not wb.sheetnames:
        ws = wb.create_sheet("导出")
        ws.append(["无导出字段"])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


@router.post("/{task_id}/export-rows")
def export_rows(task_id: int, payload: schemas.ExportRowsRequest, db: Session = Depends(get_db)):
    source_columns = [str(c).strip() for c in (payload.source_columns or []) if str(c).strip()]
    target_columns = [str(c).strip() for c in (payload.target_columns or []) if str(c).strip()]
    if not source_columns and not target_columns:
        raise HTTPException(status_code=400, detail="请至少勾选一侧要导出的字段")
    limit = min(max(int(payload.limit or 50000), 1), 100000)
    task = _get_task(db, task_id)
    allowed_source = None
    allowed_target = None
    if payload.unmatched_only:
        cache = _get_cache(task_id, payload)
        if not cache:
            raise HTTPException(status_code=400, detail="请先进行内容对比")
        allowed_source = set(cache.unmatched("source"))
        allowed_target = set(cache.unmatched("target"))

    source_conditions = _normalize_match_conditions(payload.source_conditions)
    target_conditions = _normalize_match_conditions(payload.target_conditions)
    source_logic = _normalize_condition_logic(payload.source_condition_logic)
    target_logic = _normalize_condition_logic(payload.target_condition_logic)

    source_rows: list[dict] = []
    target_rows: list[dict] = []
    source_headers = [_export_column_header(c, payload.source_table) for c in source_columns]
    target_headers = [_export_column_header(c, payload.target_table) for c in target_columns]

    def _build():
        nonlocal source_rows, target_rows
        if source_columns:
            source = _connector(task, "source")
            source_rows = _collect_export_side_rows(
                connector=source,
                database=payload.source_database,
                table=payload.source_table,
                schema_name=payload.source_schema,
                columns=source_columns,
                conditions=source_conditions,
                condition_logic=source_logic,
                joins=payload.source_joins or [],
                allowed_indices=allowed_source,
                limit=limit,
            )
        if target_columns:
            target = _connector(task, "target")
            target_rows = _collect_export_side_rows(
                connector=target,
                database=payload.target_database,
                table=payload.target_table,
                schema_name=payload.target_schema,
                columns=target_columns,
                conditions=target_conditions,
                condition_logic=target_logic,
                joins=None,
                allowed_indices=allowed_target,
                limit=limit,
            )

    try:
        _run(_build)
    except ConnectionError_ as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"导出失败: {exc}") from exc

    if not source_rows and not target_rows and (source_columns or target_columns):
        # still allow empty sheets with headers only
        pass

    buf = _build_export_workbook(source_rows, target_rows, source_headers, target_headers)
    filename = f"字段内容导出_{payload.source_table}_{payload.target_table}.xlsx"
    encoded = quote(filename)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{encoded}",
            "X-Export-Source-Count": str(len(source_rows)),
            "X-Export-Target-Count": str(len(target_rows)),
        },
    )
